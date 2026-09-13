#!/usr/bin/env python
import json, sys, zipfile
import io
from PIL import Image
from pathlib import Path
from defusedxml import ElementTree as ET
from pptx import Presentation

path = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).parents[1] / 'output' / 'advanced-authoring-probe.pptx')
assert path.is_file() and path.stat().st_size > 0, path
prs = Presentation(str(path))
assert len(prs.slides) == 3
texts = []
for slide in prs.slides:
    chunks = [sh.text for sh in slide.shapes if getattr(sh, 'has_text_frame', False)]
    for sh in slide.shapes:
        if getattr(sh, 'has_table', False):
            chunks.extend(cell.text for row in sh.table.rows for cell in row.cells)
    texts.append('\n'.join(chunks))
for expected in ('RICH text runs stay editable.', 'Raster reliability, SVG precision', 'Advanced feature matrix'):
    assert any(expected in t for t in texts), expected
assert any(getattr(sh, 'has_table', False) for sh in prs.slides[2].shapes)

with zipfile.ZipFile(path) as z:
    names = z.namelist()
    xml = {n: z.read(n) for n in names if n.endswith('.xml') or n.endswith('.rels')}
    media = [n for n in names if n.startswith('ppt/media/')]
    all_xml = b'\n'.join(xml.values())
    for name in media:
        if name.endswith('.png'):
            with Image.open(io.BytesIO(z.read(name))) as image:
                image.load()
                assert image.format == 'PNG', name
    checks = {
        'rich_text_runs': all_xml.count(b'<a:r>') >= 6,
        'character_spacing': b' spc="' in all_xml,
        'margins': any(token in all_xml for token in (b' lIns="', b' marL="')),
        'bullets': b'<a:buChar' in all_xml,
        'shape_shadow': any(int(el.get('dist', '0')) > 0 for value in xml.values() for el in ET.fromstring(value).iter('{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw')),
        'raster_image': any(n.lower().endswith('.png') for n in media),
        'svg_image': any(n.lower().endswith('.svg') for n in media),
        'merged_table_colspan': b'gridSpan="3"' in all_xml,
        'merged_table_rowspan': b'rowSpan="2"' in all_xml or b'<a:vMerge' in all_xml,
        'wide_layout': abs(prs.slide_width / prs.slide_height - 16/9) < 0.01,
    }
assert all(checks.values()), checks
report = {'ok': True, 'path': str(path.resolve()), 'slides': len(prs.slides), 'size_bytes': path.stat().st_size, 'media_parts': media, 'checks': checks, 'visuals_verified': False}
print(json.dumps(report, indent=2))
